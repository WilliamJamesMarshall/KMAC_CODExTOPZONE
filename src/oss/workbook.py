"""Read the v0.3 workbook without promoting candidates to verified facts."""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .identity import company_name_key
from .models import Capability, Evidence, Solution, Supplier, Task


REQUIRED_SHEETS = (
    "과업목록", "솔루션목록", "활동연결", "통합대조표", "세부활동근거",
    "기업목록", "설명원문", "AI스캔", "스캔근거", "범위변경", "이전스캔", "분류규칙",
)


@dataclass
class KnowledgeBase:
    tasks: dict[str, Task]
    suppliers: dict[str, Supplier]
    solutions: dict[str, Solution]
    source_chunks: dict[str, str]
    detailed_capabilities: list[Capability]
    scan_capabilities: list[Capability]
    task_aliases: dict[str, str]
    taxonomy_rules: list[dict[str, str]]
    taxonomy_changes: list[dict[str, str]]
    audit: dict[str, Any]

    def capabilities_for(self, task_id: str) -> list[Capability]:
        if task_id not in self.tasks:
            return []
        return [c for c in self.detailed_capabilities if c.task_id == task_id] + [
            c for c in self.scan_capabilities if c.task_id == task_id
        ]


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _rows(sheet: Any) -> list[tuple[int, tuple[Any, ...]]]:
    return [
        (number, row)
        for number, row in enumerate(sheet.iter_rows(min_row=6, values_only=True), start=6)
        if any(value is not None and _text(value) for value in row)
    ]


def _duplicates(values: list[str]) -> list[str]:
    return sorted(value for value, count in Counter(values).items() if count > 1)


def load_knowledge_base(path: str | Path) -> KnowledgeBase:
    """Load source rows and return an auditable, in-memory knowledge base.

    The workbook's '미검토' rows remain unreviewed. Out-of-scope task links are
    reported in audit and omitted from active candidate lookup.
    """
    workbook = load_workbook(Path(path), read_only=False, data_only=True)
    missing = sorted(set(REQUIRED_SHEETS) - set(workbook.sheetnames))
    if missing:
        raise ValueError(f"Missing required worksheets: {', '.join(missing)}")

    rows = {name: _rows(workbook[name]) for name in REQUIRED_SHEETS}
    audit: dict[str, Any] = {
        "sheet_rows": {name: len(items) for name, items in rows.items()},
        "duplicate_ids": {},
        "excluded_links": [],
        "review_status": {},
        "broken_source_refs": [],
    }

    source_chunks = {_text(r[0]): _text(r[5]) for _, r in rows["설명원문"]}

    task_rows = rows["과업목록"]
    audit["duplicate_ids"]["과업목록"] = _duplicates([_text(r[0]) for _, r in task_rows])
    tasks = {
        _text(r[0]): Task(
            task_id=_text(r[0]), name=_text(r[1]), parent_group=_text(r[2]),
            business_purpose=_text(r[4]), expected_output=_text(r[5]),
            activity_tags=_text(r[6]), boundary_rule=_text(r[7]),
            domains=_text(r[13]), activity_layer=_text(r[12]), review_status=_text(r[20]),
        )
        for _, r in task_rows
    }
    audit["review_status"]["과업목록"] = dict(Counter(t.review_status for t in tasks.values()))

    suppliers: dict[str, Supplier] = {}
    for _, r in rows["AI스캔"]:
        sid = _text(r[1])
        suppliers[sid] = Supplier(sid, _text(r[2]), None, "scan")
    for _, r in rows["기업목록"]:
        sid = _text(r[0])
        existing = suppliers.get(sid)
        suppliers[sid] = Supplier(
            sid, _text(r[1]), _text(r[5]) or _text(r[4]) or None,
            "detailed+scan" if existing else "detailed",
        )

    solution_rows = rows["솔루션목록"]
    audit["duplicate_ids"]["솔루션목록"] = _duplicates([_text(r[0]) for _, r in solution_rows])
    solutions = {
        _text(r[0]): Solution(
            solution_id=_text(r[0]), supplier_id=_text(r[10]),
            description=_text(r[2]), source_url=_text(r[11]) or None,
            review_status=_text(r[14]),
        )
        for _, r in solution_rows
    }
    audit["review_status"]["솔루션목록"] = dict(Counter(s.review_status for s in solutions.values()))

    scope_notes = {_text(r[1]): _text(r[4]) for _, r in rows["범위변경"] if _text(r[0]) == "과업"}
    detailed_capabilities: list[Capability] = []
    link_rows = rows["활동연결"]
    audit["duplicate_ids"]["활동연결"] = _duplicates([_text(r[0]) for _, r in link_rows])
    audit["review_status"]["활동연결"] = dict(Counter(_text(r[15]) for _, r in link_rows))
    for number, r in link_rows:
        link_id, solution_id, task_id, supplier_id = map(_text, (r[0], r[1], r[3], r[14]))
        if task_id not in tasks or solution_id not in solutions or supplier_id not in suppliers:
            audit["excluded_links"].append({
                "sheet": "활동연결", "row": number, "link_id": link_id,
                "task_id": task_id, "solution_id": solution_id,
                "reason": scope_notes.get(task_id) or "missing referenced entity",
            })
            continue
        for source_id in (_text(value) for value in _text(r[12]).split(";")):
            if source_id and source_id not in source_chunks:
                audit["broken_source_refs"].append({"sheet": "활동연결", "row": number, "source_id": source_id})
        source_url = solutions[solution_id].source_url or ""
        evidence = Evidence(
            source_url=source_url, text=_text(r[7]), source_kind="voucher_description",
            source_id=_text(r[12]) or link_id, mapping_method=_text(r[8]),
            review_status=_text(r[15]),
        )
        detailed_capabilities.append(Capability(
            supplier_id=supplier_id, task_id=task_id, solution_id=solution_id,
            evidence=evidence, level="detailed_unreviewed",
        ))

    scan_capabilities: list[Capability] = []
    for _, r in rows["스캔근거"]:
        task_id, supplier_id = _text(r[4]), _text(r[2])
        if task_id not in tasks or supplier_id not in suppliers:
            continue
        scan_capabilities.append(Capability(
            supplier_id=supplier_id, task_id=task_id, solution_id=None,
            evidence=Evidence(
                source_url="https://pms.ai-voucher.or.kr/supply/pool",
                text=_text(r[6]), source_kind="voucher_scan",
                source_id=_text(r[0]), mapping_method=_text(r[9]),
                review_status="미검토",
            ),
            level="scan_candidate",
        ))

    task_aliases = {
        _text(r[0]): _text(r[2]) for _, r in rows["통합대조표"] if _text(r[0]) and _text(r[2])
    }
    audit["inactive_alias_targets"] = sorted({
        target for target in task_aliases.values() if target not in tasks
    })
    for number, r in rows["세부활동근거"]:
        for source_id in (_text(value) for value in _text(r[9]).split(";")):
            if source_id and source_id not in source_chunks:
                audit["broken_source_refs"].append({"sheet": "세부활동근거", "row": number, "source_id": source_id})
    taxonomy_rules = [
        {"number": _text(r[0]), "criterion": _text(r[1]), "rule": _text(r[2]), "example": _text(r[3])}
        for _, r in rows["분류규칙"]
    ]
    taxonomy_changes = [
        {"type": _text(r[0]), "id": _text(r[1]), "name": _text(r[2]),
         "scope": _text(r[3]), "reason": _text(r[4]), "related_id": _text(r[5])}
        for _, r in rows["범위변경"]
    ]
    name_ids: defaultdict[str, set[str]] = defaultdict(set)
    for supplier in suppliers.values():
        name_ids[company_name_key(supplier.name)].add(supplier.supplier_id)
    audit["possible_supplier_duplicates"] = {
        name: sorted(ids) for name, ids in name_ids.items() if name and len(ids) > 1
    }
    audit["active_counts"] = {
        "tasks": len(tasks), "suppliers": len(suppliers), "solutions": len(solutions),
        "detailed_links": len(detailed_capabilities), "scan_evidence": len(scan_capabilities),
        "source_chunks": len(source_chunks),
    }
    workbook.close()
    return KnowledgeBase(
        tasks, suppliers, solutions, source_chunks, detailed_capabilities, scan_capabilities,
        task_aliases, taxonomy_rules, taxonomy_changes, audit,
    )
