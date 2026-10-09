"""Import v0.3 AX Task candidates only when their pool quotation still matches."""

from __future__ import annotations

import argparse
import hashlib
import json
import uuid
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .identity import company_name_key
from .pool_import import _sql_string
from .workbook import load_knowledge_base


def _compact(value: str) -> str:
    return "".join(value.split())


def _payload_sql(rows: list[dict], columns: str, types: str) -> str:
    payload = _sql_string(json.dumps(rows, ensure_ascii=False, separators=(",", ":")))
    return f"select {columns} from jsonb_to_recordset({payload}::jsonb) as x({types})"


def build_ax_import(workbook_path: Path, prepared_path: Path, output_dir: Path,
                    *, batch_size: int = 100) -> dict:
    if batch_size < 1:
        raise ValueError("batch_size must be positive")
    kb = load_knowledge_base(workbook_path)
    prepared = json.loads(prepared_path.read_text(encoding="utf-8"))
    by_name: dict[str, list[dict]] = defaultdict(list)
    for item in prepared:
        by_name[company_name_key(item["name"])].append(item)
    workbook_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    namespace = uuid.uuid5(uuid.NAMESPACE_URL, f"AX-v0.3:{workbook_hash}")
    output_dir.mkdir(parents=True, exist_ok=True)

    tasks = [
        {
            "task_id": task.task_id,
            "task_name": task.name,
            "parent_group": task.parent_group,
            "business_purpose": task.business_purpose,
            "expected_output": task.expected_output,
            "activity_tags": [part.strip() for part in task.activity_tags.split(";") if part.strip()],
            "boundary_rule": task.boundary_rule,
            "activity_layer": task.activity_layer,
            "taxonomy_version": "v0.3",
            "review_status": "unreviewed",
        }
        for task in kb.tasks.values()
    ]
    task_select = _payload_sql(
        tasks,
        "task_id, task_name, parent_group, business_purpose, expected_output, "
        "activity_tags, boundary_rule, activity_layer, taxonomy_version, review_status",
        "task_id text, task_name text, parent_group text, business_purpose text, "
        "expected_output text, activity_tags jsonb, boundary_rule text, "
        "activity_layer text, taxonomy_version text, review_status text",
    )
    (output_dir / "tasks.sql").write_text(
        "insert into public.ax_tasks (task_id, task_name, parent_group, "
        "business_purpose, expected_output, activity_tags, boundary_rule, "
        "activity_layer, taxonomy_version, review_status)\n"
        + task_select + "\non conflict (task_id) do nothing;\n", encoding="utf-8",
    )

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    scan_map: dict[str, dict] = {}
    identity_basis: dict[str, str] = {}
    for row in workbook["AI스캔"].iter_rows(min_row=6, values_only=True):
        if not row[1]:
            continue
        source_id, name, source_text = str(row[1]), str(row[2]), str(row[18] or "")
        candidates = by_name[company_name_key(name)]
        exact = [item for item in candidates
                 if _compact(item["ai_solution_text"]) == _compact(source_text)]
        if len(exact) == 1:
            scan_map[source_id] = exact[0]
            identity_basis[source_id] = "name_and_solution_text"
        elif len(candidates) == 1:
            scan_map[source_id] = candidates[0]
            identity_basis[source_id] = "unique_name_and_quote_check"
        else:
            raise ValueError(f"Unresolved scan identity: {source_id} {name}")

    rows: list[dict] = []
    skipped: dict[str, int] = defaultdict(int)
    for row in workbook["스캔근거"].iter_rows(min_row=6, values_only=True):
        if not row[0]:
            continue
        evidence_id, source_id, task_id = str(row[0]), str(row[2]), str(row[4])
        if task_id not in kb.tasks:
            skipped["scan_inactive_task"] += 1
            continue
        item = scan_map[source_id]
        quote = str(row[6] or "").strip()
        if not quote or _compact(quote) not in _compact(item["ai_solution_text"]):
            raise ValueError(f"Scan quotation missing from current pool: {evidence_id}")
        rows.append({
            "id": str(uuid.uuid5(namespace, f"scan:{evidence_id}")),
            "sply_pool_no": item["sply_pool_no"],
            "task_id": task_id,
            "evidence_text": quote,
            "evidence_type": "ai_scan",
            "evidence_strength": 1,
            "mapping_method": "legacy_ai_scan",
            "extractor_version": f"v0.3:{source_id}:{evidence_id}",
        })

    for row in workbook["활동연결"].iter_rows(min_row=6, values_only=True):
        if not row[0]:
            continue
        link_id, task_id, source_id = str(row[0]), str(row[3]), str(row[14])
        if task_id not in kb.tasks:
            skipped["detailed_inactive_task"] += 1
            continue
        supplier = kb.suppliers.get(source_id)
        candidates = by_name[company_name_key(supplier.name)] if supplier else []
        if len(candidates) != 1:
            skipped["detailed_identity_unresolved"] += 1
            continue
        item = candidates[0]
        quote = str(row[7] or "").strip()
        if not quote or _compact(quote) not in _compact(item["ai_solution_text"]):
            skipped["detailed_quote_not_in_ai_pool"] += 1
            continue
        rows.append({
            "id": str(uuid.uuid5(namespace, f"detailed:{link_id}")),
            "sply_pool_no": item["sply_pool_no"],
            "task_id": task_id,
            "evidence_text": quote,
            "evidence_type": "company_description",
            "evidence_strength": 2,
            "mapping_method": "legacy_detailed",
            "extractor_version": f"v0.3:{source_id}:{link_id}",
        })
    workbook.close()

    for start in range(0, len(rows), batch_size):
        selected = rows[start:start + batch_size]
        source = _payload_sql(
            selected,
            "id, sply_pool_no, task_id, evidence_text, evidence_type, "
            "evidence_strength, mapping_method, extractor_version",
            "id uuid, sply_pool_no bigint, task_id text, evidence_text text, "
            "evidence_type text, evidence_strength smallint, mapping_method text, "
            "extractor_version text",
        )
        sql = (
            "insert into public.capability_evidence\n"
            "  (id, supplier_id, pool_entry_id, task_id, evidence_text, evidence_type,\n"
            "   evidence_strength, verification_status, mapping_method, extractor_version)\n"
            "select x.id, p.supplier_id, p.id, x.task_id, x.evidence_text, x.evidence_type,\n"
            "       x.evidence_strength, 'candidate', x.mapping_method, x.extractor_version\n"
            f"from ({source}) x\n"
            "join public.pool_entries p on p.source_program = 'AI바우처'\n"
            "  and p.source_year = 2026 and p.sply_pool_no = x.sply_pool_no\n"
            "join public.ax_tasks t on t.task_id = x.task_id\n"
            "on conflict (id) do nothing;\n"
        )
        (output_dir / f"evidence-{start // batch_size + 1:04d}.sql").write_text(sql, encoding="utf-8")

    (output_dir / "matched-pool-ids.json").write_text(
        json.dumps(sorted({row["sply_pool_no"] for row in rows})) + "\n", encoding="utf-8",
    )

    result = {
        "workbook_sha256": workbook_hash,
        "tasks": len(tasks),
        "scan_suppliers_matched": len(scan_map),
        "scan_identity_basis": dict(sorted(
            (basis, sum(value == basis for value in identity_basis.values()))
            for basis in set(identity_basis.values())
        )),
        "scan_evidence": sum(row["mapping_method"] == "legacy_ai_scan" for row in rows),
        "detailed_evidence": sum(row["mapping_method"] == "legacy_detailed" for row in rows),
        "skipped": dict(skipped),
        "evidence_batches": (len(rows) + batch_size - 1) // batch_size,
    }
    (output_dir / "audit.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Build source-verified AX Task import SQL")
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--prepared", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--batch-size", type=int, default=100)
    args = parser.parse_args()
    print(json.dumps(build_ax_import(args.workbook, args.prepared, args.output_dir,
                                     batch_size=args.batch_size), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
