"""High-threshold, source-quoted AX Task candidates for unscanned suppliers."""

from __future__ import annotations

import argparse
import hashlib
import json
import warnings
import uuid
from collections import defaultdict
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import precision_score
from sklearn.model_selection import KFold
from sklearn.multiclass import OneVsRestClassifier
from sklearn.preprocessing import MultiLabelBinarizer

from .identity import company_name_key
from .pool_import import _sql_string


THRESHOLD = 0.65
MODEL_VERSION = "char35-logreg-v1-ax-v0.3"


def _compact(value: str) -> str:
    return "".join(value.split())


def _model() -> OneVsRestClassifier:
    return OneVsRestClassifier(
        LogisticRegression(C=3, class_weight="balanced", solver="liblinear", max_iter=300),
        n_jobs=1,
    )


def _vectorizer() -> TfidfVectorizer:
    return TfidfVectorizer(
        analyzer="char", ngram_range=(3, 5), min_df=2,
        max_features=120000, sublinear_tf=True,
    )


def _quote(text: str, vectorizer: TfidfVectorizer, estimator: LogisticRegression) -> str:
    paragraphs = [part.strip() for part in text.splitlines() if len(part.strip()) >= 30]
    if not paragraphs:
        return text[:1200].strip()
    scores = estimator.decision_function(vectorizer.transform(paragraphs))
    return paragraphs[int(np.argmax(scores))][:1200].strip()


def build_predicted_import(workbook_path: Path, prepared_path: Path, output_dir: Path,
                           *, batch_size: int = 100) -> dict:
    if batch_size < 1:
        raise ValueError("batch_size must be positive")
    prepared = json.loads(prepared_path.read_text(encoding="utf-8"))
    by_name: dict[str, list[dict]] = defaultdict(list)
    for item in prepared:
        by_name[company_name_key(item["name"])].append(item)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    scanned: list[dict] = []
    source_ids: list[str] = []
    for row in workbook["AI스캔"].iter_rows(min_row=6, values_only=True):
        if not row[1]:
            continue
        sid, name, source_text = str(row[1]), str(row[2]), str(row[18] or "")
        candidates = by_name[company_name_key(name)]
        exact = [item for item in candidates
                 if _compact(item["ai_solution_text"]) == _compact(source_text)]
        matches = exact or candidates
        if len(matches) != 1:
            raise ValueError(f"Unresolved scan identity: {sid} {name}")
        scanned.append(matches[0])
        source_ids.append(sid)
    labels: dict[str, set[str]] = defaultdict(set)
    for row in workbook["스캔근거"].iter_rows(min_row=6, values_only=True):
        if row[0]:
            labels[str(row[2])].add(str(row[4]))
    workbook.close()

    mlb = MultiLabelBinarizer()
    truth = mlb.fit_transform([labels[sid] for sid in source_ids])
    texts = [item["ai_solution_text"] for item in scanned]
    predicted = np.zeros_like(truth, dtype=bool)
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Label not .* is present in all training examples")
        for train, test in KFold(n_splits=5, shuffle=True, random_state=42).split(texts):
            vectorizer = _vectorizer()
            train_matrix = vectorizer.fit_transform([texts[i] for i in train])
            model = _model()
            model.fit(train_matrix, truth[train])
            predicted[test] = model.predict_proba(
                vectorizer.transform([texts[i] for i in test])
            ) >= THRESHOLD
    precision = precision_score(truth, predicted, average="micro", zero_division=0)
    validation_predictions = int(predicted.sum())
    if precision < 0.9 or validation_predictions < 50:
        raise ValueError("Held-out AX candidate precision gate failed")

    vectorizer = _vectorizer()
    matrix = vectorizer.fit_transform(texts)
    model = _model()
    model.fit(matrix, truth)
    scanned_pool_ids = {item["sply_pool_no"] for item in scanned}
    unscanned = [item for item in prepared if item["sply_pool_no"] not in scanned_pool_ids]
    probabilities = model.predict_proba(vectorizer.transform(
        [item["ai_solution_text"] for item in unscanned]
    ))
    namespace = uuid.uuid5(uuid.NAMESPACE_URL,
                           f"AX-predict:{MODEL_VERSION}:{hashlib.sha256(workbook_path.read_bytes()).hexdigest()}")
    rows: list[dict] = []
    for item, scores in zip(unscanned, probabilities):
        for task_index in np.flatnonzero(scores >= THRESHOLD):
            task_id = str(mlb.classes_[task_index])
            quote = _quote(item["ai_solution_text"], vectorizer, model.estimators_[task_index])
            if not quote or _compact(quote) not in _compact(item["ai_solution_text"]):
                raise ValueError(f"Prediction quotation missing: {item['sply_pool_no']} {task_id}")
            rows.append({
                "id": str(uuid.uuid5(namespace, f"{item['sply_pool_no']}:{task_id}")),
                "sply_pool_no": item["sply_pool_no"],
                "task_id": task_id,
                "evidence_text": quote,
                "confidence": round(float(scores[task_index]), 3),
            })

    output_dir.mkdir(parents=True, exist_ok=True)
    for start in range(0, len(rows), batch_size):
        payload = _sql_string(json.dumps(rows[start:start + batch_size], ensure_ascii=False,
                                         separators=(",", ":")))
        sql = (
            "insert into public.capability_evidence\n"
            "  (id, supplier_id, pool_entry_id, task_id, evidence_text, evidence_type,\n"
            "   evidence_strength, confidence, verification_status, mapping_method,\n"
            "   extractor_model, extractor_version)\n"
            "select x.id, p.supplier_id, p.id, x.task_id, x.evidence_text,\n"
            "       'company_description', 1, x.confidence, 'candidate',\n"
            f"       'text_classifier', 'sklearn-logistic-regression', '{MODEL_VERSION}'\n"
            f"from jsonb_to_recordset({payload}::jsonb) as x(\n"
            "  id uuid, sply_pool_no bigint, task_id text, evidence_text text, confidence numeric)\n"
            "join public.pool_entries p on p.source_program = 'AI바우처'\n"
            "  and p.source_year = 2026 and p.sply_pool_no = x.sply_pool_no\n"
            "join public.ax_tasks t on t.task_id = x.task_id\n"
            "on conflict (id) do nothing;\n"
        )
        (output_dir / f"predicted-{start // batch_size + 1:04d}.sql").write_text(sql, encoding="utf-8")
    (output_dir / "matched-pool-ids.json").write_text(
        json.dumps(sorted({row["sply_pool_no"] for row in rows})) + "\n", encoding="utf-8",
    )
    audit = {
        "model_version": MODEL_VERSION,
        "threshold": THRESHOLD,
        "validation": {
            "method": "5-fold out-of-fold micro precision against unreviewed v0.3 scan labels",
            "precision": round(float(precision), 4),
            "predictions": validation_predictions,
            "suppliers_with_prediction": int(predicted.any(axis=1).sum()),
        },
        "training_suppliers": len(scanned),
        "unscanned_suppliers": len(unscanned),
        "predicted_suppliers": len({row["sply_pool_no"] for row in rows}),
        "predicted_pairs": len(rows),
        "batches": (len(rows) + batch_size - 1) // batch_size,
    }
    (output_dir / "audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
                                             encoding="utf-8")
    return audit


def main() -> None:
    parser = argparse.ArgumentParser(description="Build high-threshold AX candidate SQL")
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--prepared", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--batch-size", type=int, default=100)
    args = parser.parse_args()
    print(json.dumps(build_predicted_import(args.workbook, args.prepared, args.output_dir,
                                            batch_size=args.batch_size), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
